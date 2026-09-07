"""Export violations to Excel (xlsx) using openpyxl."""
import logging
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

HEADER_FILL = PatternFill(start_color='304156', end_color='304156', fill_type='solid')
HEADER_FONT = Font(color='ffffff', bold=True, size=11)
BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin'),
)


def export_violations_to_excel(queryset, filepath: str):
    """Export a Violation queryset to an Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '违规记录'

    headers = ['规则ID', '违规项', '客户', '客服', '平台', '状态', '扣分', '时间']

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center')
        cell.border = BORDER

    for i, v in enumerate(queryset, 2):
        values = [
            v.rule.rule_id,
            v.rule.name,
            v.conversation.customer_name,
            v.employee.get_full_name() if v.employee else '',
            v.conversation.get_platform_display(),
            v.get_status_display(),
            v.adjusted_penalty or v.penalty_points,
            v.created_at.strftime('%Y-%m-%d %H:%M'),
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='center')

    # Column widths
    widths = [10, 30, 15, 15, 10, 10, 8, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(filepath)
    logger.info('Exported %d violations to %s', queryset.count(), filepath)
    return filepath